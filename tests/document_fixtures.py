import io

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def pdf_bytes(texts=("Power transformer ABC-123", "Voltage ZX-928 second page"), encrypted=False):
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    for text in texts:
        page = writer.add_blank_page(width=595, height=842)
        if text:
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 40 800 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading("电力设备资料", 0)
    doc.add_paragraph("变压器型号是 DOCX-628，以下为参数示例。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "测试电压"
    table.cell(1, 1).text = "TABLE-773"
    doc.add_paragraph("表后段落，用于检查内容顺序。")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
